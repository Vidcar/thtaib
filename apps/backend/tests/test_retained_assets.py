"""Packet03 retained asset backend slice tests."""

from __future__ import annotations

import base64
import hashlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from PIL import Image

from workbench_backend.agents.schemas import AgentEvent, AgentRun
from workbench_backend.assets.schemas import (
    RegisterVerifiedOutputRequest,
    RetainedAssetDeletionRequest,
    RetainedAssetListFilters,
    RetainedAssetReuseRequest,
    RetainedUploadRequest,
)
from workbench_backend.assets.service import RetainedAssetService
from workbench_backend.chat.schemas import ChatConversation
from workbench_backend.inference.ids import utc_now
from workbench_backend.paths import WorkbenchPaths
from workbench_backend.state.schemas import RelatedFile
from workbench_backend.state.store import ApplicationStore

from tests.support import close_workbench_sqlite


def b64(text: str | bytes) -> str:
    data = text.encode("utf-8") if isinstance(text, str) else text
    return base64.b64encode(data).decode("ascii")


class RetainedAssetServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = ApplicationStore(WorkbenchPaths(self.root))
        self.assets = RetainedAssetService(self.store)
        self.now = utc_now()

    def tearDown(self) -> None:
        close_workbench_sqlite(self.store)
        self.tmp.cleanup()

    def put_conversation(
        self,
        conversation_id: str,
        *,
        project_path: str | None = None,
        run_ids: list[str] | None = None,
    ) -> ChatConversation:
        conversation = ChatConversation(
            id=conversation_id,
            deployment_id="dep",
            area_kind="project" if project_path else "general",
            area_project_path=project_path,
            project_path=project_path,
            run_ids=run_ids or [],
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        return self.store.put_conversation(conversation)

    def test_upload_retains_immutable_session_scoped_text_and_reuses_without_tools(self) -> None:
        self.put_conversation("chat_1")
        asset = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_1",
                filename="notes.md",
                content_type="text/markdown",
                content_base64=b64("# Plan\nKeep all bytes."),
            )
        )

        self.assertEqual(asset.origin.value, "upload")
        self.assertEqual(asset.scope.value, "session")
        self.assertEqual(asset.access_scope, "session:chat_1")
        self.assertEqual(asset.size_bytes, len("# Plan\nKeep all bytes.".encode("utf-8")))
        self.assertEqual(asset.filename, "notes.md")
        self.assertIsNone(asset.mutable_reference)

        preview = self.assets.preview(asset.id, session_id="chat_1")
        self.assertEqual(preview.preview, "# Plan\nKeep all bytes.")
        content = self.assets.content(asset.id, session_id="chat_1")
        self.assertEqual(content.text, "# Plan\nKeep all bytes.")
        blocks = self.assets.current_user_content(
            RetainedAssetReuseRequest(asset_ids=[asset.id], session_id="chat_1")
        )
        self.assertEqual(len(blocks), 1)
        self.assertIn("Source retained file: notes.md", blocks[0].text)
        self.assertIn("# Plan\nKeep all bytes.", blocks[0].text)

    def test_tiny_document_values_cannot_create_unbounded_extraction_records(self) -> None:
        from workbench_backend.assets.extraction import extract_document
        self.put_conversation('chat_sections')
        for media, content in [('application/json', b'[' + b'0,' * 10000 + b'0]'), ('text/csv', b'0\n' * 10001)]:
            with self.subTest(media=media):
                with self.assertRaises(HTTPException) as rejected:
                    self.assets.retain_upload(RetainedUploadRequest(session_id='chat_sections', filename='many-rows.json' if media == 'application/json' else 'many-rows.csv', content_type=media, content_base64=b64(content)))
                self.assertEqual(rejected.exception.status_code, 413)
                self.assertIn('Split', rejected.exception.detail)
        self.assertEqual(self.assets.list_assets(), [])
        # The source labels are retained too; tiny values cannot hide huge keys
        # from the extraction budget.
        with self.assertRaises(HTTPException) as labels:
            extract_document(b'{"' + b'k' * 2000000 + b'":0}', 'application/json')
        self.assertEqual(labels.exception.status_code, 413)
        ordinary = extract_document(b'0\n' * 10000, 'text/csv')
        self.assertEqual(len(ordinary.sections), 10000)

    def test_source_reference_opens_exact_retained_version_and_range(self) -> None:
        from workbench_backend.assets.sources import SourceRangeRequest, read_source, source_url
        self.put_conversation("chat_source")
        asset = self.assets.retain_upload(RetainedUploadRequest(session_id="chat_source", filename="sales.csv", content_type="text/csv", content_base64=b64("name,total\nOrchard,42\n")))
        section = asset.extraction.sections[1]
        request = SourceRangeRequest(sha256=asset.sha256, source=section.source, extracted_line=1, start_char=2, end_char=7, session_id="chat_source")
        result = read_source(self.assets, asset.id, request)
        self.assertEqual(result.text, section.text[2:7])
        self.assertEqual(result.sha256, asset.sha256)
        self.assertIn(asset.sha256, source_url(asset, section.source, 1, 2, 7))
        for changes, status in [({"sha256": "0" * 64}, 409), ({"session_id": "unrelated"}, 403), ({"source": "invented page 9"}, 404), ({"end_char": 100000}, 422)]:
            with self.subTest(changes=changes), self.assertRaises(HTTPException) as denied:
                read_source(self.assets, asset.id, request.model_copy(update=changes))
            self.assertEqual(denied.exception.status_code, status)
        original = self.assets.content(asset.id, session_id="chat_source")
        self.assertEqual(original.text, "name,total\nOrchard,42\n")
        self.assertEqual(self.assets.list_assets()[0].extraction.sections, [])
        self.assertEqual(self.assets.store.get(asset.id)[0].extraction.sections[1].text, section.text)

    def test_attachment_admission_does_not_hydrate_other_conversations(self) -> None:
        from types import SimpleNamespace
        from workbench_backend.assets.tools import attachment_tool_for_run, validate_retained_selection
        from workbench_backend.errors import HarnessError
        own = self.put_conversation('chat_own')
        own.thread_id = 'thread_own'
        own.archived = True
        self.store.put_conversation(own)
        asset = self.assets.retain_upload(RetainedUploadRequest(session_id=own.id, filename='notes.txt', content_type='text/plain', content_base64=b64('SCOPED-RECEIPT')))
        other = self.put_conversation('chat_other')
        other.thread_id = 'thread_other'
        self.store.put_conversation(other)
        with patch.object(self.store, 'list_conversations', side_effect=AssertionError('Unrelated histories must not be hydrated')):
            validate_retained_selection(self.store, [], thread_id='thread_missing', project_path=None)
            validate_retained_selection(self.store, [asset.id], thread_id=own.thread_id, project_path=None)
            tool = attachment_tool_for_run(self.store, SimpleNamespace(retained_asset_ids=[asset.id], thread_id=own.thread_id, project_path=None))
            self.assertIn('SCOPED-RECEIPT', str(tool.invoke({'asset_id': asset.id})['excerpts']))
            with self.assertRaises(HarnessError) as denied:
                validate_retained_selection(self.store, [asset.id], thread_id=other.thread_id, project_path=None)
            self.assertEqual(denied.exception.status_code, 403)

    def test_image_upload_retains_original_and_supplies_actual_pixels_with_scoped_preview(self) -> None:
        self.put_conversation("chat_image")
        buffer = io.BytesIO()
        Image.new("RGB", (32, 24), "red").save(buffer, "PNG")
        original = buffer.getvalue()
        asset = self.assets.retain_upload(RetainedUploadRequest(
            session_id="chat_image", filename="red.png", content_type="image/png",
            content_kind="image", content_base64=b64(original),
        ))
        self.assertEqual(asset.image_width, 32)
        self.assertEqual(asset.image_height, 24)
        preview = self.assets.preview(asset.id, session_id="chat_image")
        self.assertTrue(preview.image_data_url.startswith("data:image/"))
        content = self.assets.content(asset.id, session_id="chat_image")
        self.assertEqual(base64.b64decode(content.content_base64), original)
        self.assertEqual(content.sha256, hashlib.sha256(original).hexdigest())
        blocks = self.assets.current_user_content(RetainedAssetReuseRequest(asset_ids=[asset.id], session_id="chat_image"))
        self.assertEqual(blocks[1].type, "image_url")
        self.assertEqual(base64.b64decode(blocks[1].image_url.url.split(",", 1)[1]), original)
        with self.assertRaises(HTTPException) as denied:
            self.assets.preview(asset.id, session_id="another_chat")
        self.assertEqual(denied.exception.status_code, 403)

    def test_long_document_scoped_search_retains_original_and_denies_other_files(self) -> None:
        from types import SimpleNamespace
        from workbench_backend.assets.tools import attachment_tool_for_run, validate_retained_selection
        from workbench_backend.errors import HarnessError
        conversation = self.put_conversation("chat_long")
        conversation.thread_id = "thread_long"
        self.store.put_conversation(conversation)
        text = "ordinary text " * 4000 + "NEEDLE-ORCHID-932" + " final text" * 1000
        asset = self.assets.retain_upload(RetainedUploadRequest(session_id="chat_long", filename="long.txt", content_type="text/plain", content_base64=b64(text)))
        request = RetainedAssetReuseRequest(asset_ids=[asset.id], session_id="chat_long", max_chars=2000)
        with self.assertRaises(HTTPException) as too_large:
            self.assets.current_user_content(request)
        self.assertEqual(too_large.exception.status_code, 413)
        blocks = self.assets.current_user_content(request, allow_scoped_read=True)
        self.assertIn("read_attachment", blocks[0].text)
        self.assertNotIn("NEEDLE-ORCHID", blocks[0].text)
        run = SimpleNamespace(retained_asset_ids=[asset.id], thread_id="thread_long", project_path=None)
        tool = attachment_tool_for_run(self.store, run)
        result = tool.invoke({"asset_id": asset.id, "query": "NEEDLE-ORCHID-932"})
        self.assertIn("NEEDLE-ORCHID-932", str(result["excerpts"]))
        self.assertEqual(result["sha256"], hashlib.sha256(text.encode()).hexdigest())
        first = tool.invoke({"asset_id": asset.id, "line_count": 1})
        offset = first["excerpts"][0]["next_start_char"]
        second = tool.invoke({"asset_id": asset.id, "line_count": 1, "start_char": offset})
        self.assertEqual(second["excerpts"][0]["start_char"], offset)
        self.assertEqual(first["excerpts"][0]["text"] + second["excerpts"][0]["text"], text[:24000])
        self.assertIn("not selected", tool.invoke({"asset_id": "another"}))
        with self.assertRaises(HarnessError):
            validate_retained_selection(self.store, [asset.id], thread_id="another", project_path=None)
        self.assertEqual(self.assets.content(asset.id, session_id="chat_long").text, text)

    def test_image_validation_checks_pixels_not_only_file_extension(self) -> None:
        self.put_conversation("chat_image")
        for content, media in [(b"\x89PNG\r\n\x1a\nnot pixels", "image/png"), (b"<svg onload='execute()'/>", "image/svg+xml")]:
            with self.subTest(media=media), self.assertRaises(HTTPException) as rejected:
                self.assets.retain_upload(RetainedUploadRequest(session_id="chat_image", filename="image.png",
                    content_type=media, content_kind="image", content_base64=b64(content)))
            self.assertIn(rejected.exception.status_code, {415, 422})
        buffer = io.BytesIO()
        Image.new("RGB", (10, 10), "blue").save(buffer, "JPEG")
        with self.assertRaises(HTTPException) as mismatched:
            self.assets.retain_upload(RetainedUploadRequest(session_id="chat_image", filename="fake.png",
                content_type="image/png", content_kind="image", content_base64=b64(buffer.getvalue())))
        self.assertEqual(mismatched.exception.status_code, 422)

    def test_documents_extract_source_ranges_and_preserve_originals(self) -> None:
        import docx
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        self.put_conversation("chat_docs")
        word = docx.Document()
        word.add_paragraph("Pear trees blossom in spring.")
        word.add_table(rows=1, cols=1).cell(0, 0).text = "Harvest: September"
        word_bytes = io.BytesIO()
        word.save(word_bytes)
        pdf = PdfWriter()
        page = pdf.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 30 100 Td (Orchard code: PEAR-42) Tj ET")
        page[NameObject("/Contents")] = pdf._add_object(stream)
        pdf_bytes = io.BytesIO()
        pdf.write(pdf_bytes)
        cases = [
            ("notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", word_bytes.getvalue(), "Paragraph 1", "Pear trees"),
            ("notes.pdf", "application/pdf", pdf_bytes.getvalue(), "Page 1", "PEAR-42"),
            ("notes.csv", "text/csv", b"fruit,count\npear,42\n", "Row 2", "pear | 42"),
            ("notes.json", "application/json", b'{"orchard":{"code":"PEAR-42"}}', "/orchard", "PEAR-42"),
        ]
        for name, media, original, source, expected in cases:
            with self.subTest(name=name):
                asset = self.assets.retain_upload(RetainedUploadRequest(session_id="chat_docs", filename=name,
                    content_type=media, content_base64=b64(original)))
                self.assertEqual(asset.extraction.status, "complete")
                blocks = self.assets.current_user_content(RetainedAssetReuseRequest(asset_ids=[asset.id], session_id="chat_docs"))
                self.assertIn(source, blocks[0].text)
                self.assertIn(expected, blocks[0].text)
                saved = self.assets.content(asset.id, session_id="chat_docs")
                recovered = base64.b64decode(saved.content_base64) if saved.content_base64 else saved.text.encode("utf-8")
                self.assertEqual(recovered, original)

    def test_scanned_pdf_is_retained_with_honest_no_text_state_and_not_silently_sent(self) -> None:
        from pypdf import PdfWriter
        self.put_conversation("chat_docs")
        pdf = PdfWriter()
        pdf.add_blank_page(width=300, height=200)
        buffer = io.BytesIO()
        pdf.write(buffer)
        asset = self.assets.retain_upload(RetainedUploadRequest(session_id="chat_docs", filename="scan.pdf",
            content_type="application/pdf", content_base64=b64(buffer.getvalue())))
        self.assertEqual(asset.extraction.status, "no_text")
        self.assertIn("OCR", self.assets.preview(asset.id, session_id="chat_docs").extraction.note)
        with self.assertRaises(HTTPException) as unreadable:
            self.assets.current_user_content(RetainedAssetReuseRequest(asset_ids=[asset.id], session_id="chat_docs"))
        self.assertEqual(unreadable.exception.status_code, 422)
        pdf.encrypt("secret")
        buffer = io.BytesIO()
        pdf.write(buffer)
        for original in [buffer.getvalue(), b"not a pdf"]:
            with self.subTest(encrypted=original.startswith(b"%PDF")), self.assertRaises(HTTPException) as invalid:
                self.assets.retain_upload(RetainedUploadRequest(session_id="chat_docs", filename="bad.pdf",
                    content_type="application/pdf", content_base64=b64(original)))
            self.assertEqual(invalid.exception.status_code, 422)

    def test_upload_rejects_invalid_encoding_binary_and_oversized_content(self) -> None:
        self.put_conversation("chat_1")
        with self.assertRaises(HTTPException) as bad_encoding:
            self.assets.retain_upload(
                RetainedUploadRequest(
                    session_id="chat_1",
                    filename="bad.txt",
                    content_type="text/plain",
                    content_base64=b64(b"\xff\xfe"),
                )
            )
        self.assertEqual(bad_encoding.exception.status_code, 422)

        with self.assertRaises(HTTPException) as bad_type:
            self.assets.retain_upload(
                RetainedUploadRequest(
                    session_id="chat_1",
                    filename="bad.bin",
                    content_type="application/octet-stream",
                    content_base64=b64("hello"),
                )
            )
        self.assertEqual(bad_type.exception.status_code, 415)

        with self.assertRaises(HTTPException) as too_large:
            self.assets.retain_upload(
                RetainedUploadRequest(
                    session_id="chat_1",
                    filename="large.txt",
                    content_type="text/plain",
                    content_base64=b64("x" * 1_000_001),
                )
            )
        self.assertEqual(too_large.exception.status_code, 413)

    def test_access_requires_source_session_or_same_project(self) -> None:
        project = str((self.root / "project").resolve())
        Path(project).mkdir()
        self.put_conversation("chat_project", project_path=project)
        asset = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_project",
                filename="facts.txt",
                content_type="text/plain",
                content_base64=b64("project scoped"),
            )
        )

        self.assertEqual(asset.project_path, project)
        self.assertEqual(self.assets.content(asset.id, project_path=project).text, "project scoped")
        with self.assertRaises(HTTPException) as denied:
            self.assets.content(asset.id, session_id="other_chat")
        self.assertEqual(denied.exception.status_code, 403)

    def test_upload_requires_existing_session(self) -> None:
        with self.assertRaises(HTTPException) as missing:
            self.assets.retain_upload(
                RetainedUploadRequest(
                    session_id="missing",
                    filename="notes.txt",
                    content_type="text/plain",
                    content_base64=b64("no session"),
                )
            )
        self.assertEqual(missing.exception.status_code, 404)

    def test_verified_output_reads_actual_project_file_from_failed_run_successful_tool(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("actual observed output", encoding="utf-8")
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        run = AgentRun(
            id="run_success",
            status="failed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_1",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output), "sha256": digest},
                }
            ],
            related_files=[
                RelatedFile(path=str(output), kind="written_file")
            ],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_project", project_path=str(project), run_ids=[run.id])

        branch = self.store.get_conversation("chat_project")
        branch.area_project_path = str(self.root / "historical-original-project")
        self.store.put_conversation(branch)

        asset = self.assets.register_verified_output(
            RegisterVerifiedOutputRequest(
                run_id="run_success",
                session_id="chat_project",
                source_tool_call_id="call_1",
                mutable_reference=str(output),
            )
        )

        self.assertEqual(asset.origin.value, "verified_output")
        self.assertEqual(asset.project_path, str(project))
        self.assertEqual(asset.sha256, digest)
        self.assertEqual(asset.mutable_reference, str(output))
        self.assertEqual(self.assets.content(asset.id, session_id="chat_project").text, "actual observed output")
        self.assertEqual(self.assets.preview(asset.id, session_id="chat_project").source_status, "unchanged")
        output.write_text("changed external source", encoding="utf-8")
        self.assertEqual(self.assets.preview(asset.id, session_id="chat_project").source_status, "changed")
        self.assertEqual(self.assets.content(asset.id, session_id="chat_project").text, "actual observed output")
        output.unlink()
        self.assertEqual(self.assets.content(asset.id, session_id="chat_project").source_status, "missing")
        self.assertEqual(self.assets.content(asset.id, session_id="chat_project").text, "actual observed output")

    def test_verified_output_rejects_name_fallback_hash_mismatch_and_outside_project(self) -> None:
        project = (self.root / "project").resolve()
        outside = (self.root / "outside.txt").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("actual observed output", encoding="utf-8")
        outside.write_text("outside", encoding="utf-8")
        run = AgentRun(
            id="run_strict",
            status="completed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_real",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output), "sha256": "0" * 64},
                }
            ],
            related_files=[RelatedFile(path=str(output), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_project", project_path=str(project), run_ids=[run.id])

        with self.assertRaises(HTTPException) as wrong_call:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_project",
                    source_tool_call_id="write_file",
                    mutable_reference=str(output),
                )
            )
        self.assertEqual(wrong_call.exception.status_code, 409)

        with self.assertRaises(HTTPException) as hash_mismatch:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_project",
                    source_tool_call_id="call_real",
                    mutable_reference=str(output),
                )
            )
        self.assertEqual(hash_mismatch.exception.status_code, 409)

        outside_run = run.model_copy(
            update={
                "id": "run_outside",
                "tool_invocations": [
                    {
                        "id": "call_outside",
                        "name": "write_file",
                        "status": "success",
                        "result": {"path": str(outside)},
                    }
                ],
                "related_files": [RelatedFile(path=str(outside), kind="written_file")],
            },
            deep=True,
        )
        self.store.put_run(outside_run)
        self.put_conversation("chat_outside", project_path=str(project), run_ids=[outside_run.id])
        with self.assertRaises(HTTPException) as outside_denied:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=outside_run.id,
                    session_id="chat_outside",
                    source_tool_call_id="call_outside",
                    mutable_reference=str(outside),
                )
            )
        self.assertEqual(outside_denied.exception.status_code, 403)

    def test_verified_output_rejects_successful_call_for_different_observed_path(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        wanted = project / "wanted.txt"
        other = project / "other.txt"
        wanted.write_text("wanted bytes", encoding="utf-8")
        other.write_text("other bytes", encoding="utf-8")
        run = AgentRun(
            id="run_wrong_path",
            status="completed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_other",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(other)},
                }
            ],
            related_files=[RelatedFile(path=str(wanted), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_wrong_path", project_path=str(project), run_ids=[run.id])

        with self.assertRaises(HTTPException) as rejected:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_wrong_path",
                    source_tool_call_id="call_other",
                    mutable_reference=str(wanted),
                )
            )

        self.assertEqual(rejected.exception.status_code, 409)
        self.assertIn("does not match", str(rejected.exception.detail))

    def test_verified_output_rejects_later_failed_same_file_without_success_hash(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("bytes after failed attempt", encoding="utf-8")
        run = AgentRun(
            id="run_later_failed_same_file",
            status="failed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_success",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output)},
                },
                {
                    "id": "call_failed",
                    "name": "write_file",
                    "status": "failed",
                    "result": {"path": str(output)},
                    "content": "Error: failed after partial write",
                },
            ],
            related_files=[RelatedFile(path=str(output), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_later_failed", project_path=str(project), run_ids=[run.id])

        with self.assertRaises(HTTPException) as rejected:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_later_failed",
                    source_tool_call_id="call_success",
                    mutable_reference=str(output),
                )
            )

        self.assertEqual(rejected.exception.status_code, 409)
        self.assertIn("Cannot verify current bytes", str(rejected.exception.detail))

    def test_verified_output_rejects_split_event_later_failed_same_file_without_success_hash(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("bytes after failed attempt", encoding="utf-8")
        run = AgentRun(
            id="run_split_later_failed_same_file",
            status="failed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            events=[
                AgentEvent(
                    at=utc_now(),
                    kind="tool_call",
                    detail={"id": "call_success", "name": "write_file", "args": {"path": str(output)}},
                ),
                AgentEvent(
                    at=utc_now(),
                    kind="tool_result",
                    detail={"tool_call_id": "call_success", "status": "success", "content": "Wrote file."},
                ),
                AgentEvent(
                    at=utc_now(),
                    kind="tool_call",
                    detail={"id": "call_failed", "name": "write_file", "args": {"path": str(output)}},
                ),
                AgentEvent(
                    at=utc_now(),
                    kind="tool_result",
                    detail={"tool_call_id": "call_failed", "status": "failed", "content": "Error: failed after partial write"},
                ),
            ],
            related_files=[RelatedFile(path=str(output), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_split_later_failed", project_path=str(project), run_ids=[run.id])

        with self.assertRaises(HTTPException) as rejected:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_split_later_failed",
                    source_tool_call_id="call_success",
                    mutable_reference=str(output),
                )
            )

        self.assertEqual(rejected.exception.status_code, 409)
        self.assertIn("Cannot verify current bytes", str(rejected.exception.detail))

    def test_verified_output_rejects_raw_success_when_same_call_final_result_failed(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("failed bytes", encoding="utf-8")
        run = AgentRun(
            id="run_same_call_final_failed",
            status="failed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_same",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output)},
                }
            ],
            events=[
                AgentEvent(
                    at=utc_now(),
                    kind="tool_result",
                    detail={"tool_call_id": "call_same", "status": "failed", "content": "Error: failed after write"},
                )
            ],
            related_files=[RelatedFile(path=str(output), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_same_call_failed", project_path=str(project), run_ids=[run.id])

        with self.assertRaises(HTTPException) as rejected:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_same_call_failed",
                    source_tool_call_id="call_same",
                    mutable_reference=str(output),
                )
            )

        self.assertEqual(rejected.exception.status_code, 409)
        self.assertIn("Failed or denied", str(rejected.exception.detail))

    def test_verified_output_rejects_earlier_success_when_later_success_same_file_without_hash(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("second write bytes", encoding="utf-8")
        run = AgentRun(
            id="run_two_success_same_file",
            status="completed",
            deployment_id="dep",
            task="write twice",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_first",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output)},
                },
                {
                    "id": "call_second",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output)},
                },
            ],
            related_files=[RelatedFile(path=str(output), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_two_success", project_path=str(project), run_ids=[run.id])

        with self.assertRaises(HTTPException) as rejected:
            self.assets.register_verified_output(
                RegisterVerifiedOutputRequest(
                    run_id=run.id,
                    session_id="chat_two_success",
                    source_tool_call_id="call_first",
                    mutable_reference=str(output),
                )
            )
        retained_latest = self.assets.register_verified_output(
            RegisterVerifiedOutputRequest(
                run_id=run.id,
                session_id="chat_two_success",
                source_tool_call_id="call_second",
                mutable_reference=str(output),
            )
        )

        self.assertEqual(rejected.exception.status_code, 409)
        self.assertIn("Cannot verify current bytes", str(rejected.exception.detail))
        self.assertEqual(self.assets.content(retained_latest.id, session_id="chat_two_success").text, "second write bytes")

    def test_verified_output_accepts_later_failed_same_file_when_success_hash_matches_current_bytes(self) -> None:
        project = (self.root / "project").resolve()
        project.mkdir()
        output = project / "out.txt"
        output.write_text("successful bytes", encoding="utf-8")
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        run = AgentRun(
            id="run_later_failed_hashed",
            status="failed",
            deployment_id="dep",
            task="write",
            enabled_tools=["write_file"],
            presented_tools=["write_file"],
            tool_invocations=[
                {
                    "id": "call_success",
                    "name": "write_file",
                    "status": "success",
                    "result": {"path": str(output), "sha256": digest},
                },
                {
                    "id": "call_failed",
                    "name": "write_file",
                    "status": "failed",
                    "result": {"path": str(output)},
                    "content": "Error: failed without changing retained bytes",
                },
            ],
            related_files=[RelatedFile(path=str(output), kind="written_file")],
            project_path=str(project),
            created_at=utc_now(),
            updated_at=utc_now(),
        )
        self.store.put_run(run)
        self.put_conversation("chat_later_failed_hash", project_path=str(project), run_ids=[run.id])

        asset = self.assets.register_verified_output(
            RegisterVerifiedOutputRequest(
                run_id=run.id,
                session_id="chat_later_failed_hash",
                source_tool_call_id="call_success",
                mutable_reference=str(output),
            )
        )

        self.assertEqual(asset.sha256, digest)
        self.assertEqual(self.assets.content(asset.id, session_id="chat_later_failed_hash").text, "successful bytes")

    def test_reuse_capacity_reports_actionable_outcome_without_truncation(self) -> None:
        self.put_conversation("chat_1")
        asset = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_1",
                filename="long.txt",
                content_type="text/plain",
                content_base64=b64("abcdef"),
            )
        )
        with self.assertRaises(HTTPException) as too_large:
            self.assets.current_user_content(
                RetainedAssetReuseRequest(asset_ids=[asset.id], session_id="chat_1", max_chars=10)
            )
        self.assertEqual(too_large.exception.status_code, 413)

    def test_explicit_cross_session_reuse_grants_subsequent_session_access(self) -> None:
        self.put_conversation("chat_source")
        self.put_conversation("chat_target")
        asset = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_source",
                filename="share.txt",
                content_type="text/plain",
                content_base64=b64("shared by selection"),
            )
        )

        with self.assertRaises(HTTPException):
            self.assets.content(asset.id, session_id="chat_target")
        blocks = self.assets.current_user_content(
            RetainedAssetReuseRequest(
                asset_ids=[asset.id],
                session_id="chat_target",
                allow_cross_session_reuse=True,
            )
        )
        self.assertIn("shared by selection", blocks[0].text)
        self.assertEqual(self.assets.content(asset.id, session_id="chat_target").text, "shared by selection")

    def test_session_list_includes_explicit_reuse_consumers_without_unrelated_assets(self) -> None:
        self.put_conversation("chat_source")
        self.put_conversation("chat_target")
        self.put_conversation("chat_other")
        shared = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_source",
                filename="shared.txt",
                content_type="text/plain",
                content_base64=b64("shared by selection"),
            )
        )
        unrelated = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_other",
                filename="other.txt",
                content_type="text/plain",
                content_base64=b64("other"),
            )
        )

        self.assets.current_user_content(
            RetainedAssetReuseRequest(
                asset_ids=[shared.id],
                session_id="chat_target",
                allow_cross_session_reuse=True,
            )
        )

        listed = self.assets.list_assets(RetainedAssetListFilters(session_id="chat_target"))
        self.assertEqual([asset.id for asset in listed], [shared.id])
        self.assertNotIn(unrelated.id, [asset.id for asset in listed])

        self.assets.store.mark_deleted(shared.id, utc_now())
        self.assertEqual(self.assets.list_assets(RetainedAssetListFilters(session_id="chat_target")), [])
        deleted = self.assets.list_assets(RetainedAssetListFilters(session_id="chat_target", include_deleted=True))
        self.assertEqual([asset.id for asset in deleted], [shared.id])

    def test_session_and_project_list_keeps_explicit_session_reuse(self) -> None:
        project = self.root / "target-project"
        project.mkdir()
        self.put_conversation("chat_source")
        self.put_conversation("chat_target", project_path=str(project))
        shared = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_source",
                filename="shared.txt",
                content_type="text/plain",
                content_base64=b64("shared by selection"),
            )
        )

        self.assets.current_user_content(
            RetainedAssetReuseRequest(
                asset_ids=[shared.id],
                session_id="chat_target",
                project_path=str(project),
                allow_cross_session_reuse=True,
            )
        )

        listed = self.assets.list_assets(
            RetainedAssetListFilters(session_id="chat_target", project_path=str(project))
        )
        self.assertEqual([asset.id for asset in listed], [shared.id])

    def test_deletion_preview_preserves_assets_shared_with_retained_consumers(self) -> None:
        self.put_conversation("chat_delete")
        asset = self.assets.retain_upload(
            RetainedUploadRequest(
                session_id="chat_delete",
                filename="shared.txt",
                content_type="text/plain",
                content_base64=b64("shared"),
            )
        )
        self.assets.store.add_consumer(
            asset.id,
            kind="session",
            consumer_id="chat_keeper",
            recorded_at=utc_now(),
        )

        preview = self.assets.deletion_preview(
            RetainedAssetDeletionRequest(session_ids=["chat_delete"])
        )

        self.assertEqual(preview.requested_asset_ids, [asset.id])
        self.assertEqual(preview.affected_asset_ids, [])
        self.assertEqual(preview.preserved_asset_ids, [asset.id])
        self.assertIn("chat_delete", preview.affected_sessions)
        self.assertIn("chat_keeper", preview.retained_sessions)

        deleted = self.assets.mark_deletable_assets_deleted(
            RetainedAssetDeletionRequest(session_ids=["chat_delete"])
        )
        self.assertEqual(deleted.preserved_asset_ids, [asset.id])
        self.assertEqual(self.assets.content(asset.id, session_id="chat_keeper").text, "shared")

        final = self.assets.mark_deletable_assets_deleted(
            RetainedAssetDeletionRequest(session_ids=["chat_keeper"])
        )
        self.assertEqual(final.affected_asset_ids, [asset.id])
        with self.assertRaises(HTTPException) as gone:
            self.assets.content(asset.id, session_id="chat_keeper")
        self.assertEqual(gone.exception.status_code, 410)
        loaded = self.assets.store.get(asset.id)
        self.assertIsNotNone(loaded)
        assert loaded is not None
        self.assertEqual(loaded[1], b"")


if __name__ == "__main__":
    unittest.main()
